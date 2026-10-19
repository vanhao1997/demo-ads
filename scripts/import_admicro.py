"""Read the public CPD workbook without changing it; preserve source lineage."""

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def extract(source: Path, destination: Path) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True)
    formulas = openpyxl.load_workbook(source, data_only=False)
    sheet = next(sheet for sheet in workbook if sheet.title.strip() == "CPD")
    formula_sheet = formulas[sheet.title]
    merged_origins = {}
    for region in sheet.merged_cells.ranges:
        origin = sheet.cell(region.min_row, region.min_col)
        for row in range(region.min_row, region.max_row + 1):
            for col in range(region.min_col, region.max_col + 1):
                merged_origins[(row, col)] = origin.coordinate

    rows = []
    for row_number in range(1, sheet.max_row + 1):
        cells, resolved, lineage, links, formula_values = {}, {}, {}, {}, {}
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            cell = sheet.cell(row_number, col)
            if cell.value is not None:
                cells[letter] = cell.value
            origin_address = merged_origins.get((row_number, col), cell.coordinate)
            origin = sheet[origin_address]
            if origin.value is not None:
                resolved[letter] = origin.value
                if origin_address != cell.coordinate:
                    lineage[letter] = origin_address
            if origin.hyperlink and origin.hyperlink.target:
                links[letter] = origin.hyperlink.target
            formula_cell = formula_sheet[cell.coordinate]
            if formula_cell.data_type == "f":
                formula_values[letter] = formula_cell.value
        rows.append({
            "rowNumber": row_number,
            "cells": cells,
            "resolvedCells": resolved,
            "mergedFrom": lineage,
            "hyperlinks": links,
            "formulas": formula_values,
        })
    snapshot = {
        "snapshotId": "admicro-cpd-2026-10-12",
        "sourceRefId": "E17",
        "snapshotDate": "2026-10-12",
        "effectiveAt": None,
        "sourceUrl": "https://docs.google.com/spreadsheets/d/1xB2UW-NO8rczy6_HBVBQgdem0_aSqzkEzSanuA2ntKU/edit?gid=435416766#gid=435416766",
        "workbookSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sheetName": sheet.title,
        "rowCount": sheet.max_row,
        "columnCount": sheet.max_column,
        "mergedRanges": sorted(str(region) for region in sheet.merged_cells.ranges),
        "sourceDates": ["2026-05-28", "2025-02-05", "2022-04"],
        "rows": rows,
    }
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"rows": len(rows), "columns": sheet.max_column, "sha256": snapshot["workbookSha256"], "output": str(destination)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    arguments = parser.parse_args()
    extract(arguments.source, arguments.destination)
